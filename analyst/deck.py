"""Pitch deck: Claude content blocks plus plain text for fact-checking (same rules as ic-builder/sources.py)."""

import base64
import io
from dataclasses import dataclass

from pptx import Presentation
from pypdf import PdfReader

# The API takes up to 32 MB per request; base64 adds a third, and the prompt needs room too.
MAX_PDF_BYTES = 22 * 1024 * 1024
MAX_PDF_PAGES = 600


@dataclass
class Deck:
    name: str
    text: str  # extracted text, used to check Claude's quotes
    pdf: bytes | None = None
    pages: int = 0

    @property
    def text_is_reliable(self) -> bool:
        """False for scanned PDFs, where extracted text is too thin to check quotes against."""
        return self.pdf is None or len(self.text.strip()) >= 100 * max(self.pages, 1)

    def block(self) -> dict:
        if self.pdf is not None:
            source = {"type": "base64", "media_type": "application/pdf", "data": base64.b64encode(self.pdf).decode()}
        else:
            source = {"type": "text", "media_type": "text/plain", "data": self.text or "(empty)"}
        return {"type": "document", "title": self.name, "source": source, "citations": {"enabled": True}}


def _pptx_text(data: bytes) -> str:
    prs = Presentation(io.BytesIO(data))
    out = []
    for n, slide in enumerate(prs.slides, start=1):
        out.append(f"--- Slide {n} ---")
        for sh in slide.shapes:
            out.extend(_shape_text(sh))
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame.text.strip():
            out.append(f"(Notes) {slide.notes_slide.notes_text_frame.text.strip()}")
    return "\n".join(out)


def _shape_text(sh) -> list[str]:
    if getattr(sh, "shapes", None) is not None and not getattr(sh, "has_text_frame", False):
        return [line for child in sh.shapes for line in _shape_text(child)]
    if getattr(sh, "has_table", False) and sh.has_table:
        return [" | ".join(c.text.strip() for c in row.cells) for row in sh.table.rows]
    if getattr(sh, "has_text_frame", False) and sh.has_text_frame and sh.text_frame.text.strip():
        return [sh.text_frame.text.strip()]
    return []


def load(name: str, data: bytes) -> Deck:
    lower = name.lower()
    if lower.endswith(".pdf"):
        if len(data) > MAX_PDF_BYTES:
            raise ValueError(f"{name}: PDF가 너무 큽니다 ({len(data) / 1e6:.0f}MB). 22MB 이하로 줄여 주세요.")
        reader = PdfReader(io.BytesIO(data))
        if len(reader.pages) > MAX_PDF_PAGES:
            raise ValueError(f"{name}: {len(reader.pages)}페이지입니다. {MAX_PDF_PAGES}페이지 이하로 줄여 주세요.")
        text = "\n".join((page.extract_text() or "") for page in reader.pages)
        return Deck(name, text, pdf=data, pages=len(reader.pages))
    if lower.endswith(".pptx"):
        return Deck(name, _pptx_text(data))
    raise ValueError(f"{name}: .pdf 또는 .pptx 파일만 넣을 수 있습니다. (.ppt는 PowerPoint에서 .pptx로 저장해 주세요.)")

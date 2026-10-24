import io
from datetime import datetime
from pathlib import Path

import anthropic
import pytest
from openpyxl import Workbook
from streamlit.testing.v1 import AppTest

import financials
import metrics
from financials import ItemRef, Mapping, Timeline
from metrics import Period
from research import Analyst
from test_analyst import Stream, resp
from verify import EDITED, MISSING, UNVERIFIED, VERIFIED

QUARTERS = ["Q1 2024", "Q2 2024", "Q3 2024", "Q4 2024", "Q1 2025", "Q2 2025", "Q3 2025", "Q4 2025"]


def sample_workbook() -> bytes:
    """A made-up Series C SaaS company: quarterly ARR bridge in $000s, P&L in dollars, monthly KPIs."""
    wb = Workbook()
    ws = wb.active
    ws.title = "ARR Bridge"
    ws["A1"] = "ARR Bridge ($000s)"
    ws.append([])
    ws.append(["", *QUARTERS, "Total"])
    eop, rows = 10_000, {k: [] for k in ["bop", "new", "exp", "con", "churn", "eop"]}
    for q in range(8):
        new, exp, con, churn = 1_000 + 100 * q, 500 + 50 * q, -100, -200 - 10 * q
        rows["bop"].append(eop)
        eop = eop + new + exp + con + churn
        for k, v in zip(["new", "exp", "con", "churn", "eop"], [new, exp, con, churn, eop]):
            rows[k].append(v)
    for label, k in [("Beginning ARR", "bop"), ("New Logo ARR", "new"), ("Expansion ARR", "exp"),
                     ("Contraction ARR", "con"), ("Churned ARR", "churn"), ("Ending ARR", "eop")]:
        ws.append([label, *rows[k], None])

    pl = wb.create_sheet("P&L")
    pl.append(["Income statement (USD)", *QUARTERS])
    for label, base, step in [("Revenue", 2_500_000, 200_000), ("Cost of Revenue", 600_000, 40_000),
                              ("Sales & Marketing", 1_200_000, 50_000), ("Research & Development", 900_000, 30_000),
                              ("General & Administrative", 400_000, 10_000), ("Free Cash Flow", -900_000, 80_000),
                              ("Cash Balance", 60_000_000, -900_000), ("Customers", 400, 30), ("New Customers", 40, 3)]:
        pl.append([label, *[base + step * q for q in range(8)]])
    pl.append(["Net Dollar Retention %", *[104 + q * 0.5 for q in range(8)]])

    kpi = wb.create_sheet("KPIs")
    kpi.append(["Monthly KPIs", *[datetime(2025, m, 1) for m in range(1, 7)]])
    kpi.append(["MAU", *[50_000 + 1_000 * m for m in range(6)]])
    kpi.append(["DAU", *[20_000 + 500 * m for m in range(6)]])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


SAMPLE_MAPPING = Mapping(
    currency="USD",
    timelines=[Timeline(sheet="ARR Bridge", header_row=3, first_col="B", last_col="I", granularity="quarter"),
               Timeline(sheet="P&L", header_row=1, first_col="B", last_col="I", granularity="quarter"),
               Timeline(sheet="KPIs", header_row=1, first_col="B", last_col="G", granularity="month")],
    items=[
        ItemRef(item="bop_arr", sheet="ARR Bridge", row=4, label="Beginning ARR", unit="thousands"),
        ItemRef(item="new_logo_arr", sheet="ARR Bridge", row=5, label="New Logo ARR", unit="thousands"),
        ItemRef(item="expansion_arr", sheet="ARR Bridge", row=6, label="Expansion ARR", unit="thousands"),
        ItemRef(item="contraction_arr", sheet="ARR Bridge", row=7, label="Contraction ARR", unit="thousands"),
        ItemRef(item="churned_arr", sheet="ARR Bridge", row=8, label="Churned ARR", unit="thousands"),
        ItemRef(item="arr", sheet="ARR Bridge", row=9, label="Ending ARR", unit="thousands"),
        ItemRef(item="revenue", sheet="P&L", row=2, label="Revenue", unit="units"),
        ItemRef(item="cogs", sheet="P&L", row=3, label="Cost of Revenue", unit="units"),
        ItemRef(item="sm_expense", sheet="P&L", row=4, label="Sales & Marketing", unit="units"),
        ItemRef(item="rd_expense", sheet="P&L", row=5, label="Research & Development", unit="units"),
        ItemRef(item="ga_expense", sheet="P&L", row=6, label="General & Administrative", unit="units"),
        ItemRef(item="fcf", sheet="P&L", row=7, label="Free Cash Flow", unit="units"),
        ItemRef(item="cash", sheet="P&L", row=8, label="Cash Balance", unit="units"),
        ItemRef(item="customers", sheet="P&L", row=9, label="Customers", unit="count"),
        ItemRef(item="new_customers", sheet="P&L", row=10, label="New Customers", unit="count"),
        ItemRef(item="reported_ndr", sheet="P&L", row=11, label="Net Dollar Retention %", unit="percent"),
        ItemRef(item="mau", sheet="KPIs", row=2, label="MAU", unit="count"),
        ItemRef(item="dau", sheet="KPIs", row=3, label="DAU", unit="count"),
        ItemRef(item="gmv", sheet="P&L", row=2, label="Gross Merchandise Value", unit="units"),  # wrong row
    ])


class FakeFinClient:
    def __init__(self):
        self.calls = []
        self.messages = self

    def stream(self, **kw):
        self.calls.append(kw)
        assert kw["output_format"] is Mapping
        return Stream(resp([], parsed=SAMPLE_MAPPING))


def q(year, n):
    return Period((year, n), "quarter", f"{year} Q{n}")


# --- metrics: the spec's formulas on hand-checked numbers ----------------------------------------

def one_period(**items):
    p = q(2025, 1)
    return {m.name: m.values[p] for m in metrics.compute({k: {p: v} for k, v in items.items()}, [p])}


def test_arr_bridge_metrics():
    m = one_period(bop_arr=1000, new_logo_arr=200, expansion_arr=100, contraction_arr=-20, churned_arr=-30, arr=1250,
                   sm_expense=500, fcf=-200, revenue=400, cash=2400)
    assert m["Net New ARR"] == 250 and m["Gross New ARR"] == 300
    assert m["Quick Ratio"] == pytest.approx(6.0)  # 300 / (20 + 30)
    assert m["Gross Dollar Retention"] == pytest.approx(1 - 50 / 1125)
    assert m["Net Dollar Retention"] == pytest.approx(1 + (250 - 200) / 1125)
    assert m["New Logo % of Gross New ARR"] == pytest.approx(2 / 3)
    assert m["Magic Number"] == pytest.approx(0.5)
    assert m["Burn Multiple"] == pytest.approx(0.8)  # burn 200 / net new 250
    assert m["Cash Runway"] == pytest.approx(36)  # 12 quarters of 200 burn = 36 months
    assert m["FCF Margin"] == pytest.approx(-0.5)
    assert m["Gross Churn Rate"] == pytest.approx(0.05) and m["Net Churn Rate"] == pytest.approx(-0.05)


def test_unit_economics():
    m = one_period(arr=1_200_000, bop_arr=1_000_000, churned_arr=-25_000, customers=100, new_customers=10,
                   sm_expense=50_000, revenue=300_000, cogs=60_000)
    assert m["ACV"] == 12_000 and m["Gross Margin"] == pytest.approx(0.8)
    assert m["CAC"] == 5_000
    # quarterly churn 2.5% -> 10% a year -> 10-year life; LTV = 12,000 x 0.8 x 10
    assert m["LTV"] == pytest.approx(96_000) and m["LTV / CAC"] == pytest.approx(19.2)
    assert m["CAC Payback"] == pytest.approx(5_000 / (12_000 * 0.8) * 12)


def test_missing_inputs_are_skipped_not_guessed():
    m = one_period(arr=1000, revenue=400)
    assert set(m) == {"ARR", "Revenue"}
    assert "Net Dollar Retention" not in one_period(arr=1000, bop_arr=900, churned_arr=-10)  # no expansion data


def test_yoy_and_rule_of_40_use_the_same_quarter_a_year_earlier():
    periods = [q(2024, 1), q(2024, 2), q(2025, 1)]
    series = {"arr": dict(zip(periods, [100, 110, 150])), "fcf": {q(2025, 1): 10}, "revenue": {q(2025, 1): 50}}
    rows = {m.name: m for m in metrics.compute(series, periods)}
    assert rows["YoY ARR Growth"].values == {q(2025, 1): pytest.approx(0.5)}
    assert rows["Rule of 40"].values[q(2025, 1)] == pytest.approx(0.7)  # 50% growth + 20% FCF margin


# --- reading the workbook ------------------------------------------------------------------------

@pytest.mark.parametrize("cell,gran,label", [
    ("Q1 2024", "quarter", "2024 Q1"), ("1Q24", "quarter", "2024 Q1"), ("Q3'25", "quarter", "2025 Q3"),
    ("FY2023", "year", "2023"), ("2025E", "year", "2025E"), (2022, "year", "2022"), ("Jan-24", "month", "2024-01"),
    (datetime(2024, 6, 30), "quarter", "2024 Q2"), (datetime(2025, 3, 1), "month", "2025-03"),
])
def test_period_headers(cell, gran, label):
    assert financials.parse_period(cell, gran, 2).label == label


def test_unreadable_header_keeps_its_text():
    p = financials.parse_period("Base case", "quarter", 7)
    assert p.label == "Base case" and p.key == (9999, 7)


def test_extract_reads_cells_scales_units_and_rejects_wrong_rows():
    wb = financials.load("acme.xlsx", sample_workbook())
    ex = financials.extract(wb, SAMPLE_MAPPING)
    status = {c.item: c.status for c in ex.checks}
    assert status["arr"] == VERIFIED and status["gmv"] == UNVERIFIED
    assert "gmv" not in ex.series  # its label isn't on that row, so the number isn't used
    assert ex.series["arr"][q(2024, 1)] == 11_200_000  # 11,200 in $000s
    assert ex.series["reported_ndr"][q(2024, 1)] == pytest.approx(1.04)  # written as 104
    assert {p.gran for p in ex.periods} == {"quarter", "month"}
    assert "Total" not in [p.label for p in ex.periods]


def test_rows_without_numbers_or_headers_are_missing():
    wb = financials.load("acme.xlsx", sample_workbook())
    m = Mapping(currency="USD", timelines=SAMPLE_MAPPING.timelines, items=[
        ItemRef(item="arr", sheet="ARR Bridge", row=1, label="ARR Bridge ($000s)", unit="thousands")])
    assert financials.extract(wb, m).checks[0].status == MISSING


def test_analyst_edit_skips_the_label_check():
    wb = financials.load("acme.xlsx", sample_workbook())
    m = Mapping(currency="USD", timelines=SAMPLE_MAPPING.timelines,
                items=[ItemRef(item="gmv", sheet="P&L", row=2, label="", unit="units")])
    ex = financials.extract(wb, m, edited={"gmv"})
    assert ex.checks[0].status == EDITED and ex.series["gmv"][q(2024, 1)] == 2_500_000


def test_bad_files_are_rejected():
    with pytest.raises(ValueError):
        financials.load("x.xlsx", b"not a workbook")


def test_map_items_sends_the_cells_and_item_list():
    client = FakeFinClient()
    analyst = Analyst("file:acme.xlsx", client=client)
    wb = financials.load("acme.xlsx", sample_workbook())
    assert financials.map_items(analyst, wb) is SAMPLE_MAPPING
    sent = client.calls[0]["messages"][0]["content"][0]["text"]
    assert "### Sheet: ARR Bridge" in sent and "r9: A='Ending ARR'" in sent and "- new_logo_arr:" in sent
    assert "{items}" not in sent and analyst.usage.calls == 1


# --- the tab -------------------------------------------------------------------------------------

def test_financials_tab_renders_tables_and_charts(monkeypatch):
    monkeypatch.setattr(anthropic, "Anthropic", lambda **kw: FakeFinClient())
    at = AppTest.from_file(str(Path(__file__).parent.parent / "app.py"), default_timeout=30)
    at.session_state["fin_wb"] = financials.load("acme.xlsx", sample_workbook())
    at.session_state["fin_mapping"] = SAMPLE_MAPPING
    at.run()
    assert not at.exception
    assert [s.value for s in at.subheader] == ["지표", "차트", "항목 매핑"]
    assert at.radio[0].options == ["분기별", "월별"]
    assert "Gross Dollar Retention" in at.multiselect[0].value
    assert not any(w in m.value for m in at.markdown for w in ("좋", "나쁘", "우수", "양호"))

    at.radio[0].set_value("month").run()
    assert not at.exception and at.multiselect[0].options == ["DAU", "MAU", "DAU / MAU"]

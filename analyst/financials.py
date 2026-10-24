"""Company financials (.xlsx) -> line items over time. Claude only says which row holds which item;
the numbers are read from the cells by this code, and every metric is computed in metrics.py."""

import io
import re
from dataclasses import dataclass, field
from datetime import date, datetime
from typing import Literal

from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string, get_column_letter
from pydantic import BaseModel

from metrics import COUNT, MONEY, PCT, REPORTED_ITEMS, Period
from research import Analyst, prompt
from verify import EDITED, MISSING, UNVERIFIED, VERIFIED

MAX_BYTES = 20 * 1024 * 1024
# What one sheet sends to Claude; a 5-year monthly model is ~60 columns.
MAX_COLS, MAX_ROWS = 80, 1500

# Line items Claude may map: key -> (description for Claude and the table, kind).
ITEMS: dict[str, tuple[str, str]] = {
    "arr": ("ARR (Annual Recurring Revenue), end of period", MONEY),
    "mrr": ("MRR (Monthly Recurring Revenue), end of period", MONEY),
    "bop_arr": ("Beginning-of-period ARR (opening ARR in an ARR bridge)", MONEY),
    "new_logo_arr": ("New logo ARR: ARR from new customers in the period", MONEY),
    "expansion_arr": ("Expansion / upsell ARR from existing customers in the period", MONEY),
    "contraction_arr": ("Contraction / downgrade ARR in the period", MONEY),
    "churned_arr": ("Churned / cancelled ARR in the period", MONEY),
    "revenue": ("Total revenue for the period (P&L)", MONEY),
    "cogs": ("Cost of revenue / COGS", MONEY),
    "gross_profit": ("Gross profit", MONEY),
    "sm_expense": ("Sales & marketing expense", MONEY),
    "rd_expense": ("Research & development expense", MONEY),
    "ga_expense": ("General & administrative expense", MONEY),
    "total_opex": ("Total operating expenses", MONEY),
    "operating_cf": ("Operating cash flow", MONEY),
    "capex": ("Capital expenditures", MONEY),
    "fcf": ("Free cash flow", MONEY),
    "net_burn": ("Net burn (cash burned in the period)", MONEY),
    "cash": ("Cash balance, end of period", MONEY),
    "customers": ("Number of customers, end of period", COUNT),
    "new_customers": ("Number of new customers acquired in the period", COUNT),
    "acv": ("Average contract value (ACV) per customer", MONEY),
    "gmv": ("Gross merchandise value (marketplaces)", MONEY),
    "take_rate": ("Take rate (%)", PCT),
    "arpu": ("ARPU (average revenue per user)", MONEY),
    "dau": ("Daily active users", COUNT),
    "mau": ("Monthly active users", COUNT),
    **{k: (f"{name} as already calculated in the file", kind) for k, (name, kind) in REPORTED_ITEMS.items()},
}
ITEM_KEYS = tuple(ITEMS)
UNITS = ("units", "thousands", "millions", "billions", "percent", "count")
SCALE = {"units": 1, "thousands": 1e3, "millions": 1e6, "billions": 1e9, "percent": 1, "count": 1}


# --- Workbook ------------------------------------------------------------------------------------

@dataclass
class Workbook:
    name: str
    sheets: dict[str, dict[tuple[int, int], object]]  # sheet -> (row, col) -> value, 1-based
    uncached_formulas: int = 0  # formula cells with no saved result (file never opened in Excel)

    def cell(self, sheet: str, row: int, col: int):
        return self.sheets.get(sheet, {}).get((row, col))

    def row_text(self, sheet: str, row: int) -> str:
        cells = self.sheets.get(sheet, {})
        return " ".join(str(v) for (r, _), v in sorted(cells.items()) if r == row and isinstance(v, str))

    def render(self) -> str:
        """The sheets as text Claude can read, one line per non-empty row with cell addresses."""
        out = []
        for name, cells in self.sheets.items():
            rows: dict[int, list] = {}
            for (r, c), v in sorted(cells.items()):
                if c <= MAX_COLS:
                    rows.setdefault(r, []).append((c, v))
            out.append(f"### Sheet: {name}")
            for r in sorted(rows)[:MAX_ROWS]:
                out.append(f"r{r}: " + " | ".join(f"{get_column_letter(c)}={_show(v)}" for c, v in rows[r]))
        return "\n".join(out)


def _show(v) -> str:
    if isinstance(v, datetime):
        return v.date().isoformat()
    if isinstance(v, date):
        return v.isoformat()
    if isinstance(v, float):
        return f"{v:.6g}"
    return repr(v) if isinstance(v, str) else str(v)


def load(name: str, data: bytes) -> Workbook:
    if len(data) > MAX_BYTES:
        raise ValueError(f"엑셀 파일이 너무 큽니다 ({len(data) / 1e6:.0f}MB). 20MB 이하만 받습니다.")
    try:
        values = load_workbook(io.BytesIO(data), data_only=True)
        formulas = load_workbook(io.BytesIO(data), data_only=False)
    except Exception as e:  # openpyxl raises many types for files it can't read
        raise ValueError(f"엑셀 파일을 열 수 없습니다 (.xlsx/.xlsm만 지원): {e}") from e
    sheets, uncached = {}, 0
    for ws in values.worksheets:
        if ws.sheet_state != "visible":
            continue
        fws = formulas[ws.title]
        cells = {}
        for row in ws.iter_rows():
            for c in row:
                v = c.value
                if v is None:
                    if isinstance(fws.cell(c.row, c.column).value, str) and str(fws.cell(c.row, c.column).value).startswith("="):
                        uncached += 1
                    continue
                if isinstance(v, str) and not v.strip():
                    continue
                cells[(c.row, c.column)] = v.strip() if isinstance(v, str) else v
        if cells:
            sheets[ws.title] = cells
    if not sheets:
        raise ValueError("엑셀에 읽을 수 있는 값이 없습니다.")
    return Workbook(name, sheets, uncached)


# --- Claude's mapping ----------------------------------------------------------------------------

class Timeline(BaseModel):
    sheet: str
    header_row: int  # the row whose cells name the periods
    first_col: str  # column letter of the first period
    last_col: str
    granularity: Literal["month", "quarter", "year"]


class ItemRef(BaseModel):
    item: Literal[ITEM_KEYS]
    sheet: str
    row: int
    label: str  # the row's label exactly as the file writes it
    unit: Literal[UNITS]


class Mapping(BaseModel):
    currency: str  # e.g. USD, KRW; "unknown" when the file doesn't say
    timelines: list[Timeline]
    items: list[ItemRef]


def map_items(analyst: Analyst, wb: Workbook) -> Mapping:
    items = "\n".join(f"- {k}: {desc}" for k, (desc, _) in ITEMS.items())
    material = f"<workbook name=\"{wb.name}\">\n{wb.render()}\n</workbook>"
    return analyst.extract(material, prompt("financial_analyst").replace("{items}", items), Mapping)


# --- Reading the numbers -------------------------------------------------------------------------

@dataclass
class ItemCheck:
    item: str
    sheet: str
    row: int
    label: str  # what Claude said the row is called
    unit: str
    found: str  # the row's text in the file
    status: str
    note: str = ""


@dataclass
class Extraction:
    series: dict[str, dict[Period, float]] = field(default_factory=dict)
    checks: list[ItemCheck] = field(default_factory=list)
    periods: list[Period] = field(default_factory=list)


_MONTHS = {m: i for i, m in enumerate(
    ["jan", "feb", "mar", "apr", "may", "jun", "jul", "aug", "sep", "oct", "nov", "dec"], start=1)}


def _year(s: str) -> int:
    y = int(s)
    return y + 2000 if y < 100 else y


def parse_period(v, gran: str, col: int) -> Period:
    """A period header cell -> Period. Unreadable headers keep their text and column order."""
    forecast = ""
    if isinstance(v, (datetime, date)):
        y, m = v.year, v.month
    elif isinstance(v, (int, float)) and 1990 <= v <= 2100:
        y, m = int(v), None
    else:
        s = str(v).strip()
        low = s.lower()
        forecast = "E" if re.search(r"\d\s*[efb]\b|forecast|budget|plan|전망|예상", low) else ""
        y = m = None
        if q := (re.search(r"q([1-4])\W*(?:fy|cy)?\W*(\d{4}|\d{2})(?!\d)", low)
                 or re.search(r"(?<!\d)([1-4])q\W*(\d{4}|\d{2})(?!\d)", low)):
            y, m = _year(q.group(2)), int(q.group(1)) * 3
        elif q := re.search(r"(\d{4})\W*q([1-4])", low):
            y, m = int(q.group(1)), int(q.group(2)) * 3
        elif mo := re.search(r"(jan|feb|mar|apr|may|jun|jul|aug|sep|oct|nov|dec)[a-z]*\W*(\d{4}|\d{2})(?!\d)", low):
            y, m = _year(mo.group(2)), _MONTHS[mo.group(1)]
        elif mo := re.search(r"(\d{4})\W(\d{1,2})(?!\d)", low):
            y, m = int(mo.group(1)), int(mo.group(2))
        elif mo := re.search(r"(?:fy|cy)\W*(\d{4}|\d{2})(?!\d)|(?<!\d)(\d{4})(?!\d)", low):
            y = _year(mo.group(1) or mo.group(2))
        if y is None or (m is not None and not 1 <= m <= 12):
            return Period((9999, col), gran, s)
    if gran == "month" and m:
        return Period((y, m), gran, f"{y}-{m:02d}{forecast}")
    if gran == "quarter" and m:
        q = (m - 1) // 3 + 1
        return Period((y, q), gran, f"{y} Q{q}{forecast}")
    return Period((y, 0), "year", f"{y}{forecast}") if gran == "year" else Period((9999, col), gran, str(v))


def _number(v) -> float | None:
    if isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return float(v)
    if isinstance(v, str):
        s = v.strip().replace(",", "").replace("$", "").replace("₩", "")
        neg = s.startswith("(") and s.endswith(")")
        s = s.strip("()").rstrip("%")
        try:
            return -float(s) if neg else float(s)
        except ValueError:
            return None
    return None


def _norm(s: str) -> str:
    return re.sub(r"[^0-9a-z가-힣%&]+", "", s.lower())


def extract(wb: Workbook, mapping: Mapping, edited: set[str] = frozenset()) -> Extraction:
    """Read each mapped row's numbers under its sheet's period headers.

    A row is used only when its text in the file matches the label Claude gave (or the analyst edited it),
    so a wrong row number can't slip silently into the metrics."""
    out = Extraction()
    headers: dict[str, list[tuple[int, list[tuple[int, Period]]]]] = {}  # sheet -> [(header_row, columns)]
    for t in mapping.timelines:
        try:
            a, b = column_index_from_string(t.first_col.strip().upper()), column_index_from_string(t.last_col.strip().upper())
        except ValueError:
            continue
        cols = []
        for c in range(min(a, b), max(a, b) + 1):
            v = wb.cell(t.sheet, t.header_row, c)
            if v is not None:
                cols.append((c, parse_period(v, t.granularity, c)))
        if cols:
            headers.setdefault(t.sheet, []).append((t.header_row, cols))

    seen: set[Period] = set()
    for ref in mapping.items:
        found = wb.row_text(ref.sheet, ref.row)
        check = ItemCheck(ref.item, ref.sheet, ref.row, ref.label, ref.unit, found, VERIFIED)
        out.checks.append(check)
        blocks = [h for h in headers.get(ref.sheet, []) if h[0] < ref.row]
        if ref.sheet not in wb.sheets:
            check.status, check.note = UNVERIFIED, "시트가 없습니다"
            continue
        if ref.item in edited:
            check.status = EDITED
        elif not found or not (_norm(ref.label) and (_norm(ref.label) in _norm(found) or _norm(found) in _norm(ref.label))):
            check.status, check.note = UNVERIFIED, "엑셀의 행 이름이 Claude가 말한 이름과 다릅니다"
            continue
        if not blocks:
            check.status, check.note = MISSING, "이 행 위에 기간 헤더가 없습니다"
            continue
        _, cols = max(blocks, key=lambda h: h[0])  # the nearest period header above the row
        values = {p: _number(wb.cell(ref.sheet, ref.row, c)) for c, p in cols}
        values = {p: v for p, v in values.items() if v is not None}
        if not values:
            check.status, check.note = MISSING, "기간 열에 숫자가 없습니다"
            continue
        if ref.unit == "percent" or ITEMS[ref.item][1] == PCT:
            if any(abs(v) > 1.5 for v in values.values()):  # written as 95 rather than 0.95
                values = {p: v / 100 for p, v in values.items()}
        else:
            values = {p: v * SCALE[ref.unit] for p, v in values.items()}
        series = out.series.setdefault(ref.item, {})
        for p, v in values.items():
            series.setdefault(p, v)  # the first row mapped to an item wins
            seen.add(p)
    out.periods = sorted(seen)
    return out

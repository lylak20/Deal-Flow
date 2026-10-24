"""1. Template: break the old IC deck into text slots and visuals that the filler can address by id."""

from dataclasses import dataclass, field

from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Emu

MARK = "[확인 필요]"


@dataclass
class Slot:
    id: str  # "<shape path>" or "<shape path>/r<row>c<col>" for a table cell
    kind: str  # "text" | "table_cell"
    shape_name: str
    paragraphs: list[str]  # non-empty paragraphs, in order
    top_pct: int
    left_pct: int
    width_pct: int

    @property
    def text(self) -> str:
        return "\n".join(self.paragraphs)


@dataclass
class Visual:
    id: str
    kind: str  # "picture" | "chart"
    shape_name: str
    description: str
    image: tuple[str, bytes] | None  # (content type, blob) for pictures
    box: tuple[int, int, int, int]  # left, top, width, height in EMU (top-level shape)


@dataclass
class Table:
    id: str
    rows: list[list[str | None]]  # cell text by row; None for a cell merged into another


@dataclass
class SlideInfo:
    number: int  # 1-based
    slots: list[Slot] = field(default_factory=list)
    visuals: list[Visual] = field(default_factory=list)
    tables: list[Table] = field(default_factory=list)


def _walk(shapes, prefix="", top=None):
    """Yield (id, shape, top-level ancestor) for every shape, descending into groups."""
    for sh in shapes:
        sid = f"{prefix}{sh.shape_id}"
        anchor = top or sh
        if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from _walk(sh.shapes, sid + ".", anchor)
        else:
            yield sid, sh, anchor


def _box(sh) -> tuple[int, int, int, int]:
    return (int(sh.left or 0), int(sh.top or 0), int(sh.width or 0), int(sh.height or 0))


def text_frames(slide) -> dict:
    """Slot id -> text_frame, for every text box and table cell on the slide."""
    frames = {}
    for sid, sh, _ in _walk(slide.shapes):
        if getattr(sh, "has_text_frame", False) and sh.has_text_frame:
            frames[sid] = sh.text_frame
        elif getattr(sh, "has_table", False) and sh.has_table:
            for r, row in enumerate(sh.table.rows):
                for c, cell in enumerate(row.cells):
                    if cell.is_spanned:
                        continue
                    frames[f"{sid}/r{r}c{c}"] = cell.text_frame
    return frames


def visual_shapes(slide) -> dict:
    """Visual id -> (shape, top-level ancestor) for pictures and charts."""
    out = {}
    for sid, sh, anchor in _walk(slide.shapes):
        if getattr(sh, "has_chart", False) and sh.has_chart:
            out[sid] = (sh, anchor)
        elif _is_picture(sh):
            out[sid] = (sh, anchor)
    return out


def _is_picture(sh) -> bool:
    try:
        sh.image  # Picture and filled picture placeholders
        return True
    except (AttributeError, ValueError):
        return False


def _chart_description(chart) -> str:
    parts = [f"chart type {chart.chart_type}"]
    try:
        if chart.has_title and chart.chart_title.has_text_frame:
            parts.append(f"title '{chart.chart_title.text_frame.text}'")
        plot = chart.plots[0]
        parts.append(f"categories {list(plot.categories)[:12]}")
        parts.append(f"series {[s.name for s in plot.series][:8]}")
    except (IndexError, KeyError, ValueError, AttributeError):
        pass
    return ", ".join(parts)


def shapes_by_id(slide) -> dict:
    """Shape id (as used in slot ids) -> shape, including shapes inside groups."""
    return {sid: sh for sid, sh, _ in _walk(slide.shapes)}


def tables(slide) -> dict:
    """Table shape id -> table."""
    return {sid: sh.table for sid, sh, _ in _walk(slide.shapes) if getattr(sh, "has_table", False) and sh.has_table}


def extract(prs) -> list[SlideInfo]:
    return [extract_slide(slide, n, prs.slide_width, prs.slide_height) for n, slide in enumerate(prs.slides, start=1)]


def extract_slide(slide, number: int, width, height) -> SlideInfo:
    width, height = width or Emu(1), height or Emu(1)
    info = SlideInfo(number)
    anchors = {sid: anchor for sid, _, anchor in _walk(slide.shapes)}
    names = {sid: sh.name for sid, sh, _ in _walk(slide.shapes)}
    for slot_id, tf in text_frames(slide).items():
        paragraphs = [p.text.strip() for p in tf.paragraphs if p.text.strip()]
        is_cell = "/" in slot_id
        if not paragraphs and not is_cell:  # empty table cells stay addressable so they can be filled
            continue
        shape_id = slot_id.split("/")[0]
        left, top, w, _ = _box(anchors[shape_id])
        info.slots.append(Slot(
            id=slot_id,
            kind="table_cell" if is_cell else "text",
            shape_name=names[shape_id],
            paragraphs=paragraphs,
            top_pct=round(100 * top / height),
            left_pct=round(100 * left / width),
            width_pct=round(100 * w / width),
        ))
    info.slots.sort(key=lambda s: (s.top_pct, s.left_pct))
    for tid, table in tables(slide).items():
        info.tables.append(Table(tid, [
            [None if cell.is_spanned else cell.text.strip() for cell in row.cells] for row in table.rows
        ]))
    for vid, (sh, anchor) in visual_shapes(slide).items():
        if getattr(sh, "has_chart", False) and sh.has_chart:
            info.visuals.append(Visual(vid, "chart", sh.name, _chart_description(sh.chart), None, _box(anchor)))
        else:
            image = sh.image
            info.visuals.append(Visual(
                vid, "picture", sh.name, f"picture '{image.filename}'",
                (image.content_type, image.blob), _box(anchor),
            ))
    return info


def outline(slides: list[SlideInfo]) -> str:
    """The whole template as text, so each per-slide call sees the deck it belongs to."""
    lines = []
    for s in slides:
        lines.append(f"=== Slide {s.number} ===")
        for slot in s.slots:
            if slot.kind == "text":
                lines.append(f"[{slot.id}] {slot.text}")
        for t in s.tables:
            lines.append(f"[{t.id}] table:")
            lines.extend(f"  r{r}: " + " | ".join("" if c is None else c for c in row) for r, row in enumerate(t.rows))
        for v in s.visuals:
            lines.append(f"[{v.id}] <{v.kind}: {v.description}>")
    return "\n".join(lines)

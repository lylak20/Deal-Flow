"""5-7. Writer: put the new text into the template's own runs, trim tables, add slides, and look for leftovers."""

import copy
import re
from dataclasses import dataclass, field

from pptx.dml.color import RGBColor
from pptx.oxml.ns import qn
from pptx.text.text import _Run
from pptx.util import Pt

from fill import term_pattern
from template import MARK, shapes_by_id, tables, text_frames, visual_shapes

RED = RGBColor(0xE0, 0x00, 0x00)
VISUAL_LABEL = "[차트/이미지 업데이트 필요]"
NEW_SLIDE_TAG = "[신규 슬라이드]"
_R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
_INLINE = {qn("a:r"), qn("a:br"), qn("a:fld")}


def set_paragraph_text(p, text: str) -> None:
    """Replace a paragraph's text, keeping its formatting.

    Leading runs that the new text still starts with (a bold "Fund Size:" label) stay as they are;
    the rest of the text goes into the next run, so it keeps that run's font, size and colour.
    """
    text = text.replace("\n", " ")
    runs = p.runs
    if not runs:
        p.add_run().text = text
        return
    kept, rest = 0, text
    while kept < len(runs) - 1 and runs[kept].text and rest.startswith(runs[kept].text):
        rest = rest[len(runs[kept].text):]
        kept += 1
    target = runs[kept]._r
    target_seen = False
    for child in list(p._p):
        if child is target:
            target_seen = True
        elif target_seen and child.tag in _INLINE:
            p._p.remove(child)
        elif not target_seen and child.tag in (qn("a:br"), qn("a:fld")):
            p._p.remove(child)
    runs[kept].text = rest
    _mark_red(p)


def _mark_red(p) -> None:
    """Split each run around MARK so the marker alone is red."""
    for run in list(p.runs):
        parts = [x for x in re.split(f"({re.escape(MARK)})", run.text) if x]
        if MARK not in parts:
            continue
        for part in parts:
            r = copy.deepcopy(run._r)
            run._r.addprevious(r)
            new = _Run(r, p)
            new.text = part
            if part == MARK:
                new.font.color.rgb = RED
        p._p.remove(run._r)


def set_frame_paragraphs(tf, texts: list[str]) -> None:
    """Write `texts` over the frame's non-empty paragraphs, cloning or removing paragraphs as needed."""
    filled = [p for p in tf.paragraphs if p.text.strip()] or [tf.paragraphs[0]]
    if not texts:
        texts = [MARK]
    for i, text in enumerate(texts):
        if i < len(filled):
            set_paragraph_text(filled[i], text)
        else:
            last = filled[-1]._p
            clone = copy.deepcopy(last)
            last.addnext(clone)
            para = next(p for p in tf.paragraphs if p._p is clone)
            filled.append(para)
            set_paragraph_text(para, text)
    for extra in filled[len(texts):]:
        extra._p.getparent().remove(extra._p)


def _red_label(slide, left: int, top: int, width: int, text: str) -> None:
    tb = slide.shapes.add_textbox(left, top, width, Pt(20))
    tb.fill.solid()
    tb.fill.fore_color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    run = tb.text_frame.paragraphs[0].add_run()
    run.text = text
    run.font.size = Pt(10)
    run.font.bold = True
    run.font.color.rgb = RED


def label_visual(slide, box: tuple[int, int, int, int]) -> None:
    left, top, width, _ = box
    _red_label(slide, left, top, max(width, Pt(160)), VISUAL_LABEL)


def tag_new_slide(slide, slide_width: int) -> None:
    width = Pt(110)
    _red_label(slide, slide_width - width - Pt(8), Pt(6), width, NEW_SLIDE_TAG)


# --- Tables -------------------------------------------------------------------------------------


def _set_row_span(tc, span: int) -> None:
    if span > 1:
        tc.set("rowSpan", str(span))
    elif "rowSpan" in tc.attrib:
        del tc.attrib["rowSpan"]


def _remove_row(tbl, r: int) -> None:
    """Remove row r, keeping vertical merges that start in it or run through it valid."""
    trs = tbl.tr_lst
    tr = trs[r]
    for c, tc in enumerate(tr.tc_lst):
        span = int(tc.get("rowSpan", "1"))
        if span > 1 and r + 1 < len(trs):
            # The merged cell starts here: hand it (text and all) to the next row.
            below = trs[r + 1].tc_lst[c]
            moved = copy.deepcopy(tc)
            _set_row_span(moved, span - 1)
            below.addprevious(moved)
            below.getparent().remove(below)
        elif tc.get("vMerge"):
            for above in range(r - 1, -1, -1):
                origin = trs[above].tc_lst[c]
                if not origin.get("vMerge"):
                    if "rowSpan" in origin.attrib:
                        _set_row_span(origin, int(origin.get("rowSpan")) - 1)
                    break
    tbl.remove(tr)


def delete_rows(slide, frame, rows: list[int]) -> int:
    """Delete table rows (never the header row 0), and the shapes sitting on them, such as company logos.

    Shapes on rows below a deleted row move up with the table. Returns the number of rows deleted.
    """
    tbl = frame.table._tbl
    rows = sorted({r for r in rows if 0 < r < len(tbl.tr_lst)})
    if not rows:
        return 0
    bands, y = [], int(frame.top or 0)
    for tr in tbl.tr_lst:
        bands.append((y, y + int(tr.h)))
        y += int(tr.h)
    table_top, table_bottom = bands[0][0], bands[-1][1]
    left, right = int(frame.left or 0), int(frame.left or 0) + int(frame.width or 0)
    for r in reversed(rows):
        _remove_row(tbl, r)

    for sh in list(slide.shapes):
        if sh.shape_id == frame.shape_id or sh.top is None or sh.height is None:
            continue
        cx, cy = int(sh.left or 0) + int(sh.width or 0) // 2, int(sh.top) + int(sh.height) // 2
        if not (left <= cx <= right and table_top <= cy <= table_bottom):
            continue
        if any(bands[r][0] <= cy < bands[r][1] for r in rows):
            sh._element.getparent().remove(sh._element)
        else:
            shift = sum(bands[r][1] - bands[r][0] for r in rows if bands[r][1] <= cy)
            if shift:
                sh.top = int(sh.top) - shift
    return len(rows)


def delete_empty_columns(frame) -> None:
    """Drop columns with no text in any row (a new slide copied a wider table than its topic needs).

    The freed width goes to the remaining columns, so the table keeps its size.
    """
    tbl = frame.table._tbl
    trs, cols = tbl.tr_lst, tbl.tblGrid.gridCol_lst

    def empty(c: int) -> bool:
        for tr in trs:
            tc = tr.tc_lst[c]
            if tc.get("hMerge") or int(tc.get("gridSpan", "1")) > 1 or "".join(tc.itertext()).strip():
                return False
        return True

    drop = [c for c in range(len(cols)) if empty(c)]
    if not drop or len(drop) == len(cols):
        return
    freed = sum(int(cols[c].get("w")) for c in drop)
    for c in reversed(drop):
        for tr in trs:
            tr.remove(tr.tc_lst[c])
        tbl.tblGrid.remove(cols[c])
    kept = tbl.tblGrid.gridCol_lst
    total = sum(int(col.get("w")) for col in kept)
    for col in kept:
        col.set("w", str(int(col.get("w")) + freed * int(col.get("w")) // total))


# --- Slides -------------------------------------------------------------------------------------


def _needs_relationships(element) -> bool:
    """True if the shape points at another part (image, chart, link), which a plain copy would break."""
    return any(k.startswith(f"{{{_R_NS}}}") for el in element.iter() for k in el.attrib)


def duplicate_slide(prs, src):
    """Append a copy of `src` with the same layout, keeping text boxes, tables and plain shapes.

    Pictures, charts and anything else linked to another part are left out.
    """
    new = prs.slides.add_slide(src.slide_layout)
    for ph in list(new.shapes):
        ph._element.getparent().remove(ph._element)
    tree = new.shapes._spTree
    for sh in src.shapes:
        if _needs_relationships(sh._element):
            continue
        tree.insert_element_before(copy.deepcopy(sh._element), "p:extLst")
    return new


def move_last_slides(prs, count: int, before: int) -> None:
    """Move the last `count` slides so they start at 0-based position `before`."""
    ids = prs.slides._sldIdLst
    moving = list(ids)[-count:]
    for el in moving:
        ids.remove(el)
    for i, el in enumerate(moving):
        ids.insert(before + i, el)


@dataclass
class SlideEdits:
    texts: dict[str, list[str]] = field(default_factory=dict)  # slot id -> new paragraphs
    flagged: list[str] = field(default_factory=list)  # visual ids to label
    delete_rows: dict[str, list[int]] = field(default_factory=dict)  # table shape id -> rows
    remove: list[str] = field(default_factory=list)  # text shape ids to delete
    trim_columns: bool = False  # drop empty table columns (new slides)


def apply(prs, edits: dict[int, SlideEdits]) -> int:
    """Apply each slide's edits (keyed by 1-based slide number). Returns the number of table rows deleted."""
    deleted = 0
    for number, slide in enumerate(prs.slides, start=1):
        e = edits.get(number)
        if e is None:
            continue
        frames = text_frames(slide)
        for slot_id, texts in e.texts.items():
            if slot_id in frames:
                set_frame_paragraphs(frames[slot_id], texts)
        shapes = shapes_by_id(slide)
        for table_id, rows in e.delete_rows.items():
            if table_id in shapes and getattr(shapes[table_id], "has_table", False):
                deleted += delete_rows(slide, shapes[table_id], rows)
        if e.trim_columns:
            for table_id in tables(slide):
                delete_empty_columns(shapes[table_id])
        for shape_id in e.remove:
            if shape_id in shapes:
                el = shapes[shape_id]._element
                el.getparent().remove(el)
        visuals = visual_shapes(slide)
        for vid in e.flagged:
            if vid in visuals:
                label_visual(slide, _anchor_box(visuals[vid][1]))
    return deleted


def _anchor_box(sh) -> tuple[int, int, int, int]:
    return (int(sh.left or 0), int(sh.top or 0), int(sh.width or 0), int(sh.height or 0))


def leftovers(prs, terms: list[str]) -> list[tuple[int, str, str]]:
    """(slide, term, where) for each old-fund term still in the deck's text, tables or notes."""
    pattern = term_pattern(terms)
    if pattern is None:
        return []
    canonical = {t.strip().lower(): t.strip() for t in terms}
    found = []
    for number, slide in enumerate(prs.slides, start=1):
        places = [("본문", tf.text) for tf in text_frames(slide).values()]
        if slide.has_notes_slide:
            places.append(("노트", slide.notes_slide.notes_text_frame.text))
        for where, text in places:
            for m in pattern.finditer(text):
                found.append((number, canonical.get(m.group(0).lower(), m.group(0)), where))
    return sorted(set(found))

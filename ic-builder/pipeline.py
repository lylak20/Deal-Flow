"""Template + sources -> new IC deck."""

import io
from concurrent.futures import ThreadPoolExecutor, as_completed
from dataclasses import dataclass, field

from pptx import Presentation

import template
import writer
from fill import Checker, Filler, SlideDecision
from sources import Source, check_limits

MAX_NEW_SLIDES = 5


@dataclass
class Result:
    pptx: bytes
    old_fund: str
    new_fund: str
    filled: int = 0  # paragraphs written from the sources
    marked: dict[int, int] = field(default_factory=dict)  # slide -> paragraphs left as MARK
    demoted: int = 0  # paragraphs Claude wrote that failed the fact check
    flagged_visuals: dict[int, int] = field(default_factory=dict)
    deleted_rows: int = 0
    new_slides: list[tuple[int, str]] = field(default_factory=list)  # (slide, topic)
    leftovers: list[tuple[int, str, str]] = field(default_factory=list)


def run(template_bytes: bytes, sources: list[Source], on_progress=lambda done, total, stage="fill": None,
        client=None, workers: int = 4, max_new_slides: int = MAX_NEW_SLIDES) -> Result:
    check_limits(sources)
    prs = Presentation(io.BytesIO(template_bytes))
    slides = template.extract(prs)
    filler = Filler(sources, template.outline(slides), client=client)

    # One call first so the shared prefix (sources + template) is cached before the slides fan out.
    analysis = filler.analyze()
    old_terms = [analysis.old_fund_name, *analysis.old_fund_terms]
    checker = Checker(sources, old_terms)
    result = Result(b"", analysis.old_fund_name, analysis.new_fund_name)

    # Phase 1: rewrite the template's own slides.
    jobs = [(s, None) for s in slides if s.slots or s.visuals]
    decisions = _fill_all(filler, jobs, len(slides), workers, lambda d, t: on_progress(d, t, "fill"))
    edits = {s.number: _edits(s, decisions[s.number], checker, result, is_new=False)
             for s in slides if s.number in decisions}

    # Phase 2: new slides for what the sources cover but the template has no place for.
    new_jobs = []
    if max_new_slides:
        candidates = [s.number for s in slides[1:-1] if s.slots]
        proposals = filler.propose_slides(_rewritten_text(slides, edits), candidates, max_new_slides)
        for p in proposals:
            dup = writer.duplicate_slide(prs, prs.slides[p.base_slide - 1])
            info = template.extract_slide(dup, len(prs.slides), prs.slide_width, prs.slide_height)
            new_jobs.append((info, p.topic))
        if new_jobs:
            new_decisions = _fill_all(filler, new_jobs, len(prs.slides), workers,
                                      lambda d, t: on_progress(d, t, "new"))
            for info, _ in new_jobs:
                edits[info.number] = _edits(info, new_decisions[info.number], checker, result, is_new=True)

    result.deleted_rows = writer.apply(prs, edits)
    new_count = len(new_jobs)

    # New slides go before a closing page ("End of Document"), otherwise at the end.
    original = len(slides)
    closing = analysis.last_slide_is_closing and original > 1
    if new_count and closing:
        writer.move_last_slides(prs, new_count, original - 1)

    def final_number(n: int) -> int:
        if not closing or not new_count:
            return n
        if n == original:
            return original + new_count
        return n - 1 if n > original else n

    for info, topic in new_jobs:
        number = final_number(info.number)
        writer.tag_new_slide(prs.slides[number - 1], prs.slide_width)
        result.new_slides.append((number, topic))
    result.marked = {final_number(n): c for n, c in result.marked.items()}
    result.flagged_visuals = {final_number(n): c for n, c in result.flagged_visuals.items()}

    result.leftovers = writer.leftovers(prs, old_terms)
    out = io.BytesIO()
    prs.save(out)
    result.pptx = out.getvalue()
    return result


def _fill_all(filler: Filler, jobs, total: int, workers: int, progress) -> dict[int, SlideDecision]:
    decisions = {}
    progress(0, len(jobs))
    with ThreadPoolExecutor(max_workers=workers) as pool:
        futures = {pool.submit(filler.fill_slide, info, total, topic): info for info, topic in jobs}
        for done, fut in enumerate(as_completed(futures), start=1):
            decisions[futures[fut].number] = fut.result()
            progress(done, len(jobs))
    return decisions


def _edits(info: template.SlideInfo, decision: SlideDecision, checker: Checker, result: Result,
           is_new: bool) -> writer.SlideEdits:
    """Fact-check a slide's decisions and turn them into edits, counting what was filled and marked."""
    e = writer.SlideEdits(trim_columns=is_new)
    for d in decision.delete_rows:
        if d.row > 0:
            e.delete_rows.setdefault(d.table_id, []).append(d.row)
    deleted = {(t, r) for t, rows in e.delete_rows.items() for r in rows}

    by_id = {d.slot_id: d for d in decision.slots}
    for slot in info.slots:
        if slot.kind == "table_cell":
            table_id, cell = slot.id.split("/")
            if (table_id, int(cell[1:cell.index("c")])) in deleted:
                continue
        d = by_id.get(slot.id)
        kept = d is None or d.action == "keep"
        if not kept and d.action == "remove" and is_new and slot.kind == "text":
            e.remove.append(slot.id)
            continue
        if kept:
            texts, demoted = checker.check_kept(slot)
        elif d.action == "remove" or not d.paragraphs:
            texts, demoted = ([] if slot.kind == "table_cell" and is_new else [template.MARK]), 0
        else:
            texts, demoted = checker.check(slot, d.paragraphs)
        # Old-fund names Claude left in place, even in text it chose to keep, become markers.
        scrubbed = checker.scrub(texts)
        if kept and scrubbed == texts:
            continue
        texts = scrubbed or [""]  # [""]: an emptied cell on a new slide
        e.texts[slot.id] = texts
        result.demoted += demoted
        marks = sum(template.MARK in t for t in texts)
        if marks:
            result.marked[info.number] = result.marked.get(info.number, 0) + marks
        result.filled += sum(bool(t) and template.MARK not in t for t in texts)

    known = {v.id for v in info.visuals}
    e.flagged = list(dict.fromkeys(d.visual_id for d in decision.visuals if d.action == "flag" and d.visual_id in known))
    if e.flagged:
        result.flagged_visuals[info.number] = len(e.flagged)
    return e


def _rewritten_text(slides: list[template.SlideInfo], edits: dict[int, writer.SlideEdits]) -> str:
    """The deck's text after phase 1, for the new-slide proposal call."""
    lines = []
    for s in slides:
        e = edits.get(s.number, writer.SlideEdits())
        deleted = {(t, r) for t, rows in e.delete_rows.items() for r in rows}
        lines.append(f"=== Slide {s.number} ===")
        for slot in s.slots:
            if slot.kind == "table_cell":
                table_id, cell = slot.id.split("/")
                if (table_id, int(cell[1:cell.index("c")])) in deleted:
                    continue
            text = "\n".join(e.texts.get(slot.id, slot.paragraphs))
            if text.strip():
                lines.append(f"[{slot.id}] {text}")
    return "\n".join(lines)

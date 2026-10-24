"""2. Sources: the new fund's materials, as Claude content blocks plus plain text for fact-checking."""

import base64
import io
from dataclasses import dataclass

from pptx import Presentation
from pypdf import PdfReader

# The API takes up to 32 MB per request; base64 adds a third, and the prompt needs room too.
MAX_PDF_BYTES = 22 * 1024 * 1024
MAX_PDF_PAGES = 600


@dataclass
class Source:
    name: str
    kind: str  # "pdf" | "pptx" | "notes"
    text: str  # extracted text, used to check Claude's quotes
    pdf: bytes | None = None
    pages: int = 0

    @property
    def text_is_reliable(self) -> bool:
        """False for scanned PDFs, where extracted text is too thin to check quotes against."""
        return self.kind != "pdf" or len(self.text.strip()) >= 100 * max(self.pages, 1)


def _pptx_text(data: bytes) -> str:
    prs = Presentation(io.BytesIO(data))
    out = []
    for n, slide in enumerate(prs.slides, start=1):
        out.append(f"--- Slide {n} ---")
        for sh in slide.shapes:
            out.extend(_shape_text(sh))
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame.text.strip():
            out.append(f"(Notes) {slide.notes_slide.notes_text_frame.text.strip()}")
    return "\n".join(out)


def _shape_text(sh) -> list[str]:
    if getattr(sh, "shapes", None) is not None and not getattr(sh, "has_text_frame", False):
        return [line for child in sh.shapes for line in _shape_text(child)]
    if getattr(sh, "has_table", False) and sh.has_table:
        return [" | ".join(c.text.strip() for c in row.cells) for row in sh.table.rows]
    if getattr(sh, "has_chart", False) and sh.has_chart:
        lines = []
        try:
            for plot in sh.chart.plots:
                cats = list(plot.categories)
                for s in plot.series:
                    lines.append(f"(Chart) {s.name}: " + ", ".join(f"{c}={v}" for c, v in zip(cats, s.values)))
        except (KeyError, ValueError, AttributeError):
            pass
        return lines
    if getattr(sh, "has_text_frame", False) and sh.has_text_frame and sh.text_frame.text.strip():
        return [sh.text_frame.text.strip()]
    return []


def load(name: str, data: bytes) -> Source:
    lower = name.lower()
    if lower.endswith(".pdf"):
        if len(data) > MAX_PDF_BYTES:
            raise ValueError(f"{name}: PDF가 너무 큽니다 ({len(data) / 1e6:.0f}MB). 22MB 이하로 나눠 주세요.")
        reader = PdfReader(io.BytesIO(data))
        text = "\n".join((page.extract_text() or "") for page in reader.pages)
        return Source(name, "pdf", text, pdf=data, pages=len(reader.pages))
    if lower.endswith(".pptx"):
        return Source(name, "pptx", _pptx_text(data))
    raise ValueError(f"{name}: .pdf 또는 .pptx 파일만 넣을 수 있습니다.")


def check_limits(sources: list[Source]) -> None:
    total = sum(len(s.pdf) for s in sources if s.pdf)
    if total > MAX_PDF_BYTES:
        raise ValueError(f"PDF 합계가 {total / 1e6:.0f}MB입니다. 22MB 이하로 줄여 주세요.")
    pages = sum(s.pages for s in sources)
    if pages > MAX_PDF_PAGES:
        raise ValueError(f"PDF 합계 {pages}페이지입니다. {MAX_PDF_PAGES}페이지 이하로 줄여 주세요.")


def content_blocks(sources: list[Source]) -> list[dict]:
    blocks = []
    for s in sources:
        if s.pdf is not None:
            blocks.append({
                "type": "document",
                "title": s.name,
                "source": {"type": "base64", "media_type": "application/pdf", "data": base64.b64encode(s.pdf).decode()},
            })
        else:
            blocks.append({
                "type": "document",
                "title": s.name,
                "source": {"type": "text", "media_type": "text/plain", "data": s.text or "(empty)"},
            })
    return blocks
